import os, json, threading
from pathlib import Path
from datetime import datetime, date
from typing import Optional
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("UNISPHERE_XLSX", BASE_DIR / "unisphere_database.xlsx"))
DIST_DIR = BASE_DIR / "dist"
LOCK = threading.RLock()

app = FastAPI(title="UniSphere University Portal API - Excel Edition")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

wb = None
sheets = {}

def load_data():
    global wb, sheets
    with LOCK:
        wb = load_workbook(DB_PATH)
        sheets = {name: wb[name] for name in wb.sheetnames}

def save_data():
    with LOCK:
        wb.save(DB_PATH)

def rows(sheet_name):
    ws = sheets[sheet_name]
    values = list(ws.iter_rows(values_only=True))
    headers = list(values[0])
    return [dict(zip(headers, r)) for r in values[1:]]

def append_row(sheet_name, data):
    ws=sheets[sheet_name]
    headers=[c.value for c in ws[1]]
    ws.append([data.get(h) for h in headers])
    return dict(zip(headers,[data.get(h) for h in headers]))

def next_id(sheet_name):
    ids=[r.get("id") for r in rows(sheet_name) if isinstance(r.get("id"),int)]
    return max(ids,default=0)+1

def find_user(user_id):
    for r in rows("users"):
        if int(r["id"]) == int(user_id):
            return r
    return None

def update_row(sheet_name, row_id, updates):
    ws=sheets[sheet_name]
    headers=[c.value for c in ws[1]]
    id_col=headers.index("id")+1
    for rn in range(2, ws.max_row+1):
        if ws.cell(rn,id_col).value == row_id:
            for k,v in updates.items():
                if k in headers:
                    ws.cell(rn,headers.index(k)+1).value=v
            return True
    return False

def delete_row(sheet_name,row_id):
    ws=sheets[sheet_name]
    headers=[c.value for c in ws[1]]
    id_col=headers.index("id")+1
    for rn in range(2,ws.max_row+1):
        if ws.cell(rn,id_col).value == row_id:
            ws.delete_rows(rn,1)
            return True
    return False

class LoginRequest(BaseModel):
    email: str
    password: str
    role: Optional[str]=None

class AssignmentRequest(BaseModel):
    title: str
    subject: str
    description: str=""
    due_date: str
    max_marks: int=20
    created_by: int

class SubmissionRequest(BaseModel):
    assignment_id: int
    student_id: int
    file_name: str="submission.pdf"

class GradeRequest(BaseModel):
    marks: float
    feedback: str=""

class AttendanceRequest(BaseModel):
    student_id: int
    subject: str
    present: bool
    date: Optional[str]=None

class NoticeRequest(BaseModel):
    title: str
    body: str
    category: str="General"
    author: str="Academic Office"

class UserRequest(BaseModel):
    name: str
    email: str
    password: str="1234"
    role: str
    department: str
    class_year: str="TY"
    phone: str="+91 90000 00000"

class FeeRequest(BaseModel):
    student_id: int
    amount: float

class ResultRequest(BaseModel):
    student_id: int
    subject: str
    internal: float
    external: float
    total: float
    grade: str
    published: bool=True

class FacultyAttendanceRequest(BaseModel):
    faculty_id: int
    date: Optional[str]=None
    present: bool

@app.on_event("startup")
def startup():
    load_data()

@app.get("/api/health")
def health():
    return {"status":"ok","database":"unisphere_database.xlsx","server_time":datetime.now().isoformat()}

@app.post("/api/login")
def login(payload: LoginRequest):
    with LOCK:
        login_email = payload.email.strip().lower()
        aliases = {
            "student0001@gmail.com": "nikki@unisphere.edu",
            "student0002@gmail.com": "aarav@gmail.com",
            "student0003@gmail.com": "meera@gmail.com",
            "student0004@gmail.com": "kabir@gmail.com",
            "faculty01@gmail.com": "priya@unisphere.edu",
            "faculty02@gmail.com": "rajesh@unisphere.edu",
        }
        login_email = aliases.get(login_email, login_email)
        for u in rows("users"):
            if str(u["email"]).strip().lower()==login_email and str(u["password"])==payload.password:
                if payload.role and str(u["role"]).lower()!=payload.role.lower():
                    raise HTTPException(401,"Role mismatch")
                return u
    raise HTTPException(401,"Invalid email or password")

@app.get("/api/state")
def state():
    with LOCK:
        users=rows("users")
        assignments=rows("assignments")
        submissions=rows("submissions")
        notices=rows("notices")
        results=rows("results")
        placements=rows("placements")
        timetable=rows("timetable")
        fa=rows("faculty_attendance")
        user_map={u["id"]:u for u in users}
        a_map={a["id"]:a for a in assignments}
        out_sub=[]
        for s in submissions:
            s=dict(s)
            u=user_map.get(s["student_id"],{})
            a=a_map.get(s["assignment_id"],{})
            s["student_name"]=u.get("name")
            s["roll_no"]=u.get("roll_no")
            s["assignment_title"]=a.get("title")
            s["subject"]=a.get("subject")
            s["max_marks"]=a.get("max_marks")
            out_sub.append(s)
        out_res=[]
        for r in results:
            r=dict(r); u=user_map.get(r["student_id"],{})
            r["student_name"]=u.get("name"); r["roll_no"]=u.get("roll_no")
            out_res.append(r)
        out_assign=[]
        for a in assignments:
            a=dict(a); u=user_map.get(a["created_by"],{})
            a["faculty_name"]=u.get("name")
            out_assign.append(a)
        out_fa=[]
        for x in fa:
            x=dict(x); u=user_map.get(x["faculty_id"],{})
            x["faculty_name"]=u.get("name"); x["department"]=u.get("department"); x["employee_id"]=u.get("employee_id")
            out_fa.append(x)
        return {"users":users,"assignments":out_assign,"submissions":out_sub,"notices":notices,"results":out_res,"placements":placements,"timetable":timetable,"facultyAttendance":out_fa}

@app.get("/api/attendance/{student_id}")
def attendance(student_id:int):
    with LOCK:
        agg={}
        for r in rows("attendance"):
            if int(r["student_id"])==student_id:
                key=r["subject"]; agg.setdefault(key,[0,0])
                agg[key][0]+=1; agg[key][1]+=int(r["present"] or 0)
        return [{"subject":k,"total_classes":v[0],"present_classes":v[1],"percentage":round(v[1]/v[0]*100,1) if v[0] else 0} for k,v in agg.items()]

@app.post("/api/attendance")
def mark_attendance(payload:AttendanceRequest):
    dt=payload.date or date.today().isoformat()
    with LOCK:
        ws=sheets["attendance"]; headers=[c.value for c in ws[1]]
        for rn in range(2,ws.max_row+1):
            if ws.cell(rn,2).value==payload.student_id and ws.cell(rn,3).value==payload.subject and str(ws.cell(rn,4).value)==dt:
                ws.cell(rn,5).value=1 if payload.present else 0; save_data()
                return {"success":True}
        append_row("attendance",{"id":next_id("attendance"),"student_id":payload.student_id,"subject":payload.subject,"date":dt,"present":1 if payload.present else 0})
        save_data(); return {"success":True}

@app.post("/api/faculty-attendance")
def mark_faculty_attendance(payload:FacultyAttendanceRequest):
    dt=payload.date or date.today().isoformat()
    with LOCK:
        ws=sheets["faculty_attendance"]
        for rn in range(2,ws.max_row+1):
            if ws.cell(rn,2).value==payload.faculty_id and str(ws.cell(rn,3).value)==dt:
                ws.cell(rn,4).value=1 if payload.present else 0; save_data(); return {"success":True}
        append_row("faculty_attendance",{"id":next_id("faculty_attendance"),"faculty_id":payload.faculty_id,"date":dt,"present":1 if payload.present else 0})
        save_data(); return {"success":True}

@app.get("/api/faculty-attendance")
def get_faculty_attendance():
    with LOCK:
        return rows("faculty_attendance")

@app.post("/api/results")
def publish_result(payload:ResultRequest):
    with LOCK:
        for r in rows("results"):
            if int(r["student_id"])==payload.student_id and r["subject"]==payload.subject:
                update_row("results",int(r["id"]),payload.model_dump()); save_data()
                return next(x for x in rows("results") if int(x["id"])==int(r["id"]))
        rid=next_id("results")
        append_row("results",{"id":rid,**payload.model_dump()})
        save_data()
        return next(x for x in rows("results") if int(x["id"])==rid)

@app.post("/api/assignments")
def create_assignment(payload:AssignmentRequest):
    with LOCK:
        item={"id":next_id("assignments"),**payload.model_dump(),"created_at":datetime.now().isoformat()}
        append_row("assignments",item); save_data()
        return item

@app.post("/api/submissions")
def submit_assignment(payload:SubmissionRequest):
    with LOCK:
        for s in rows("submissions"):
            if int(s["assignment_id"])==payload.assignment_id and int(s["student_id"])==payload.student_id:
                raise HTTPException(400,"Assignment already submitted")
        if not any(int(a["id"])==payload.assignment_id for a in rows("assignments")):
            raise HTTPException(404,"Assignment not found")
        item={"id":next_id("submissions"),**payload.model_dump(),"submitted_at":datetime.now().isoformat(),"status":"Submitted","marks":None,"feedback":""}
        append_row("submissions",item); save_data(); return {"success":True}

@app.put("/api/submissions/{submission_id}/grade")
def grade_submission(submission_id:int,payload:GradeRequest):
    with LOCK:
        if not update_row("submissions",submission_id,payload.model_dump()|{"status":"Graded"}):
            raise HTTPException(404,"Submission not found")
        save_data(); return {"success":True}

@app.post("/api/notices")
def create_notice(payload:NoticeRequest):
    with LOCK:
        item={"id":next_id("notices"),**payload.model_dump(),"created_at":datetime.now().isoformat()}
        append_row("notices",item); save_data(); return item

@app.post("/api/users")
def create_user(payload:UserRequest):
    role=payload.role.lower()
    if role not in ("student","faculty"):
        raise HTTPException(400,"Only student or faculty accounts can be created")
    with LOCK:
        if any(str(u["email"]).lower()==payload.email.strip().lower() for u in rows("users")):
            raise HTTPException(400,"User with this email already exists")
        users=rows("users")
        if role=="student":
            n=sum(1 for u in users if u["role"]=="student")+1
            dept_code="".join(x[0] for x in payload.department.split())[:4].upper()
            code=f"{dept_code}-{n:04d}"
            emp=""
        else:
            n=sum(1 for u in users if u["role"]=="faculty")+1
            code=""
            emp=f"FAC-{n:03d}"
        item={"id":next_id("users"),"name":payload.name,"email":payload.email.strip().lower(),"password":payload.password,"role":role,"department":payload.department,"class_year":payload.class_year,"roll_no":code if role=="student" else "","employee_id":emp,"phone":payload.phone,"status":"Active"}
        append_row("users",item)
        if role=="student":
            append_row("fees",{"id":next_id("fees"),"student_id":item["id"],"semester_fee":85000,"paid":0})
        save_data(); return item

@app.delete("/api/users/{user_id}")
def delete_user(user_id:int):
    with LOCK:
        if not delete_row("users",user_id): raise HTTPException(404,"User not found")
        save_data(); return {"success":True}

@app.get("/api/fees")
def get_all_fees():
    with LOCK:
        users={u["id"]:u for u in rows("users")}
        return [{**f,"student_name":users.get(f["student_id"],{}).get("name"),"roll_no":users.get(f["student_id"],{}).get("roll_no"),"department":users.get(f["student_id"],{}).get("department")} for f in rows("fees")]

@app.get("/api/fees/{student_id}")
def get_fee(student_id:int):
    with LOCK:
        for f in rows("fees"):
            if int(f["student_id"])==student_id: return f
        return None

@app.post("/api/fees/pay")
def pay_fee(payload:FeeRequest):
    # Kept for admin/system workflows. The current student UI intentionally does not expose payment.
    with LOCK:
        for f in rows("fees"):
            if int(f["student_id"])==payload.student_id:
                new_paid=min(float(f["semester_fee"]),float(f["paid"])+float(payload.amount))
                update_row("fees",int(f["id"]),{"paid":new_paid}); save_data()
                return {"success":True,"paid":new_paid}
        raise HTTPException(404,"Fee record not found")

@app.get("/api/admin/analytics")
def admin_analytics():
    with LOCK:
        users=rows("users"); fees=rows("fees")
        students=sum(u["role"]=="student" for u in users); faculty=sum(u["role"]=="faculty" for u in users)
        total_fee=sum(float(f["semester_fee"] or 0) for f in fees); paid_fee=sum(float(f["paid"] or 0) for f in fees)
        return {"students":students,"faculty":faculty,"total_fee":total_fee,"paid_fee":paid_fee,"fee_collection_percentage":round(paid_fee/total_fee*100,1) if total_fee else 0,"assignments":len(rows("assignments")),"submissions":len(rows("submissions")),"notices":len(rows("notices"))}

load_data()

if DIST_DIR.exists():
    @app.get("/{path:path}")
    async def serve_frontend(path:str):
        requested=DIST_DIR/path
        if requested.exists() and requested.is_file(): return FileResponse(requested)
        return FileResponse(DIST_DIR/"index.html")

if __name__=="__main__":
    import uvicorn
    uvicorn.run("server:app",host="0.0.0.0",port=int(os.getenv("PORT","8000")))
